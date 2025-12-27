# -*- coding: utf-8 -*-

from odoo import models, api, _
import logging
import json

_logger = logging.getLogger(__name__)


class MigrationSchemaReader(models.AbstractModel):
    _name = 'migration.schema.reader'
    _description = 'Lector de Esquemas de Base de Datos'

    @api.model
    def read_schema(self, connection):
        """
        Leer esquema de la base de datos origen.
        Retorna lista de tablas con sus columnas.
        """
        db_type = connection.db_type

        if db_type == 'postgresql':
            return self._read_postgresql_schema(connection)
        elif db_type == 'mysql':
            return self._read_mysql_schema(connection)
        elif db_type == 'mssql':
            return self._read_mssql_schema(connection)
        elif db_type == 'oracle':
            return self._read_oracle_schema(connection)
        elif db_type == 'odoo':
            return self._read_odoo_schema(connection)
        elif db_type in ('csv', 'excel', 'json'):
            return self._read_file_schema(connection)
        else:
            raise ValueError(f'Tipo de BD no soportado: {db_type}')

    @api.model
    def analyze_relationships(self, tables):
        """
        Analizar y detectar relaciones entre tablas.
        Retorna un grafo de dependencias.
        """
        relationships = []
        table_names = {t['name'].lower(): t['name'] for t in tables}

        for table in tables:
            table_name = table['name']

            for col in table.get('columns', []):
                col_name = col.get('name', '').lower()

                # Detectar FK explícitas
                if col.get('is_fk') and col.get('fk_table'):
                    relationships.append({
                        'from_table': table_name,
                        'from_column': col['name'],
                        'to_table': col['fk_table'],
                        'to_column': 'id',
                        'type': 'foreign_key',
                        'confidence': 100,
                    })

                # Detectar FK implícitas por convención de nombres
                elif col_name.endswith('_id') and col_name != 'id':
                    potential_table = col_name[:-3]  # Quitar _id

                    # Buscar tabla que coincida
                    for table_key, real_name in table_names.items():
                        if potential_table in table_key or table_key in potential_table:
                            relationships.append({
                                'from_table': table_name,
                                'from_column': col['name'],
                                'to_table': real_name,
                                'to_column': 'id',
                                'type': 'inferred',
                                'confidence': 80,
                            })
                            break

        return relationships

    @api.model
    def get_migration_order(self, tables, relationships):
        """
        Determinar orden óptimo de migración basado en dependencias.
        Usa ordenamiento topológico.
        """
        from collections import defaultdict, deque

        # Construir grafo de dependencias
        graph = defaultdict(list)
        in_degree = defaultdict(int)
        all_tables = {t['name'] for t in tables}

        for rel in relationships:
            from_table = rel['from_table']
            to_table = rel['to_table']

            if from_table in all_tables and to_table in all_tables:
                graph[to_table].append(from_table)
                in_degree[from_table] += 1

        # Inicializar in_degree para todas las tablas
        for table in all_tables:
            if table not in in_degree:
                in_degree[table] = 0

        # Ordenamiento topológico (Kahn's algorithm)
        queue = deque([t for t in all_tables if in_degree[t] == 0])
        order = []
        priority = 1

        while queue:
            table = queue.popleft()
            order.append({
                'table': table,
                'priority': priority,
                'dependencies': [r['to_table'] for r in relationships if r['from_table'] == table]
            })
            priority += 1

            for dependent in graph[table]:
                in_degree[dependent] -= 1
                if in_degree[dependent] == 0:
                    queue.append(dependent)

        # Si hay ciclos, agregar las tablas restantes al final
        remaining = all_tables - {o['table'] for o in order}
        for table in remaining:
            order.append({
                'table': table,
                'priority': priority,
                'dependencies': [],
                'has_cycle': True
            })
            priority += 1

        return order

    @api.model
    def generate_schema_visualization(self, tables, relationships):
        """
        Generar datos para visualización del esquema.
        Retorna estructura compatible con diagramas ER.
        """
        nodes = []
        edges = []

        # Clasificar tablas por tipo
        table_categories = self._categorize_tables(tables)

        for table in tables:
            table_name = table['name']
            category = table_categories.get(table_name, 'other')

            # Determinar color por categoría
            colors = {
                'master': '#28a745',      # Verde - Datos maestros
                'transaction': '#007bff', # Azul - Transacciones
                'config': '#6c757d',      # Gris - Configuración
                'relation': '#ffc107',    # Amarillo - Tablas de relación
                'other': '#17a2b8',       # Cyan - Otros
            }

            columns_preview = table.get('columns', [])[:5]

            nodes.append({
                'id': table_name,
                'label': table_name,
                'category': category,
                'color': colors.get(category, '#17a2b8'),
                'row_count': table.get('row_count', 0),
                'column_count': len(table.get('columns', [])),
                'columns': [{
                    'name': c['name'],
                    'type': c.get('type', 'unknown'),
                    'is_pk': c.get('is_pk', False),
                    'is_fk': c.get('is_fk', False),
                } for c in columns_preview],
                'pk': table.get('pk_columns', []),
            })

        for rel in relationships:
            edges.append({
                'from': rel['from_table'],
                'to': rel['to_table'],
                'label': rel['from_column'],
                'type': rel['type'],
                'confidence': rel.get('confidence', 100),
                'dashes': rel['type'] == 'inferred',  # Línea punteada para relaciones inferidas
            })

        return {
            'nodes': nodes,
            'edges': edges,
            'stats': {
                'total_tables': len(tables),
                'total_relationships': len(relationships),
                'master_tables': len([n for n in nodes if n['category'] == 'master']),
                'transaction_tables': len([n for n in nodes if n['category'] == 'transaction']),
            }
        }

    def _categorize_tables(self, tables):
        """Categorizar tablas por su probable función"""
        categories = {}

        master_keywords = ['customer', 'client', 'partner', 'vendor', 'supplier',
                          'product', 'item', 'employee', 'user', 'country',
                          'currency', 'category', 'tax', 'account', 'warehouse']

        transaction_keywords = ['order', 'invoice', 'payment', 'shipment', 'picking',
                               'move', 'entry', 'journal', 'transaction', 'receipt']

        config_keywords = ['config', 'setting', 'parameter', 'option', 'sequence']

        relation_keywords = ['rel', 'link', 'mapping', 'association']

        for table in tables:
            name_lower = table['name'].lower()

            # Detectar tablas de relación (M2M)
            if '_rel' in name_lower or name_lower.count('_') >= 2:
                columns = table.get('columns', [])
                fk_count = sum(1 for c in columns if c.get('is_fk') or c['name'].endswith('_id'))
                if fk_count >= 2 and len(columns) <= 4:
                    categories[table['name']] = 'relation'
                    continue

            # Categorizar por keywords
            for kw in master_keywords:
                if kw in name_lower:
                    categories[table['name']] = 'master'
                    break
            else:
                for kw in transaction_keywords:
                    if kw in name_lower:
                        categories[table['name']] = 'transaction'
                        break
                else:
                    for kw in config_keywords:
                        if kw in name_lower:
                            categories[table['name']] = 'config'
                            break
                    else:
                        categories[table['name']] = 'other'

        return categories

    def _read_postgresql_schema(self, connection):
        """Leer esquema de PostgreSQL"""
        import psycopg2

        tables = []

        try:
            conn = psycopg2.connect(
                host=connection.host,
                port=connection.port,
                database=connection.database,
                user=connection.username,
                password=connection.password,
            )
            cursor = conn.cursor()

            # Obtener tablas (excluyendo tablas de sistema)
            cursor.execute("""
                SELECT
                    table_schema,
                    table_name
                FROM information_schema.tables
                WHERE table_schema NOT IN ('information_schema', 'pg_catalog')
                AND table_type = 'BASE TABLE'
                ORDER BY table_schema, table_name
            """)

            for schema, table_name in cursor.fetchall():
                # Obtener columnas
                cursor.execute("""
                    SELECT
                        column_name,
                        data_type,
                        is_nullable,
                        column_default,
                        character_maximum_length
                    FROM information_schema.columns
                    WHERE table_schema = %s AND table_name = %s
                    ORDER BY ordinal_position
                """, (schema, table_name))

                columns = []
                for col in cursor.fetchall():
                    columns.append({
                        'name': col[0],
                        'type': col[1],
                        'nullable': col[2] == 'YES',
                        'default': col[3],
                        'max_length': col[4],
                    })

                # Obtener PKs
                cursor.execute("""
                    SELECT kcu.column_name
                    FROM information_schema.table_constraints tc
                    JOIN information_schema.key_column_usage kcu
                        ON tc.constraint_name = kcu.constraint_name
                    WHERE tc.table_schema = %s
                        AND tc.table_name = %s
                        AND tc.constraint_type = 'PRIMARY KEY'
                """, (schema, table_name))
                pk_columns = [row[0] for row in cursor.fetchall()]

                for col in columns:
                    col['is_pk'] = col['name'] in pk_columns

                # Obtener FKs
                cursor.execute("""
                    SELECT
                        kcu.column_name,
                        ccu.table_name AS foreign_table_name
                    FROM information_schema.table_constraints tc
                    JOIN information_schema.key_column_usage kcu
                        ON tc.constraint_name = kcu.constraint_name
                    JOIN information_schema.constraint_column_usage ccu
                        ON ccu.constraint_name = tc.constraint_name
                    WHERE tc.table_schema = %s
                        AND tc.table_name = %s
                        AND tc.constraint_type = 'FOREIGN KEY'
                """, (schema, table_name))
                fk_info = {row[0]: row[1] for row in cursor.fetchall()}

                for col in columns:
                    if col['name'] in fk_info:
                        col['is_fk'] = True
                        col['fk_table'] = fk_info[col['name']]
                    else:
                        col['is_fk'] = False

                # Obtener conteo de registros
                try:
                    cursor.execute(f'SELECT COUNT(*) FROM "{schema}"."{table_name}"')
                    row_count = cursor.fetchone()[0]
                except:
                    row_count = 0

                tables.append({
                    'name': table_name,
                    'schema': schema,
                    'columns': columns,
                    'row_count': row_count,
                    'pk_columns': pk_columns,
                })

            conn.close()

        except Exception as e:
            _logger.error(f'Error leyendo esquema PostgreSQL: {str(e)}')
            raise

        return tables

    def _read_mysql_schema(self, connection):
        """Leer esquema de MySQL"""
        import pymysql

        tables = []

        try:
            conn = pymysql.connect(
                host=connection.host,
                port=connection.port,
                database=connection.database,
                user=connection.username,
                password=connection.password,
            )
            cursor = conn.cursor()

            # Obtener tablas
            cursor.execute(f"""
                SELECT table_name
                FROM information_schema.tables
                WHERE table_schema = '{connection.database}'
                AND table_type = 'BASE TABLE'
            """)

            for (table_name,) in cursor.fetchall():
                # Obtener columnas
                cursor.execute(f"""
                    SELECT
                        column_name,
                        data_type,
                        is_nullable,
                        column_default,
                        character_maximum_length,
                        column_key
                    FROM information_schema.columns
                    WHERE table_schema = '{connection.database}'
                    AND table_name = '{table_name}'
                    ORDER BY ordinal_position
                """)

                columns = []
                pk_columns = []
                for col in cursor.fetchall():
                    col_data = {
                        'name': col[0],
                        'type': col[1],
                        'nullable': col[2] == 'YES',
                        'default': col[3],
                        'max_length': col[4],
                        'is_pk': col[5] == 'PRI',
                        'is_fk': col[5] == 'MUL',
                    }
                    columns.append(col_data)
                    if col_data['is_pk']:
                        pk_columns.append(col[0])

                # Conteo de registros
                try:
                    cursor.execute(f'SELECT COUNT(*) FROM `{table_name}`')
                    row_count = cursor.fetchone()[0]
                except:
                    row_count = 0

                tables.append({
                    'name': table_name,
                    'schema': connection.database,
                    'columns': columns,
                    'row_count': row_count,
                    'pk_columns': pk_columns,
                })

            conn.close()

        except Exception as e:
            _logger.error(f'Error leyendo esquema MySQL: {str(e)}')
            raise

        return tables

    def _read_odoo_schema(self, connection):
        """Leer esquema de otra instancia de Odoo"""
        import xmlrpc.client

        tables = []

        try:
            common = xmlrpc.client.ServerProxy(f'{connection.odoo_url}/xmlrpc/2/common')
            uid = common.authenticate(
                connection.database,
                connection.username,
                connection.password,
                {}
            )

            models = xmlrpc.client.ServerProxy(f'{connection.odoo_url}/xmlrpc/2/object')

            # Obtener lista de modelos
            model_list = models.execute_kw(
                connection.database, uid, connection.password,
                'ir.model', 'search_read',
                [[('transient', '=', False)]],
                {'fields': ['model', 'name']}
            )

            for model_info in model_list:
                model_name = model_info['model']

                # Obtener campos del modelo
                try:
                    fields_info = models.execute_kw(
                        connection.database, uid, connection.password,
                        model_name, 'fields_get',
                        [],
                        {'attributes': ['string', 'type', 'required', 'relation']}
                    )

                    columns = []
                    for field_name, field_data in fields_info.items():
                        columns.append({
                            'name': field_name,
                            'type': field_data.get('type'),
                            'nullable': not field_data.get('required', False),
                            'is_pk': field_name == 'id',
                            'is_fk': field_data.get('type') in ('many2one', 'many2many', 'one2many'),
                            'fk_table': field_data.get('relation'),
                        })

                    # Conteo de registros
                    try:
                        row_count = models.execute_kw(
                            connection.database, uid, connection.password,
                            model_name, 'search_count', [[]]
                        )
                    except:
                        row_count = 0

                    tables.append({
                        'name': model_name,
                        'schema': 'odoo',
                        'columns': columns,
                        'row_count': row_count,
                        'pk_columns': ['id'],
                        'is_odoo_model': True,
                        'odoo_model_name': model_info['name'],
                    })

                except Exception as e:
                    _logger.warning(f'Error leyendo modelo {model_name}: {str(e)}')
                    continue

        except Exception as e:
            _logger.error(f'Error leyendo esquema Odoo: {str(e)}')
            raise

        return tables

    def _read_file_schema(self, connection):
        """Leer esquema de archivos (CSV, Excel, JSON)"""
        tables = []

        for attachment in connection.file_ids:
            file_name = attachment.name
            file_data = attachment.raw

            if connection.db_type == 'csv':
                columns = self._parse_csv_schema(file_data)
            elif connection.db_type == 'excel':
                columns = self._parse_excel_schema(file_data)
            elif connection.db_type == 'json':
                columns = self._parse_json_schema(file_data)
            else:
                continue

            tables.append({
                'name': file_name.rsplit('.', 1)[0],
                'schema': 'file',
                'columns': columns,
                'row_count': 0,  # Se calculará al leer el archivo
                'pk_columns': [],
                'file_id': attachment.id,
            })

        return tables

    def _parse_csv_schema(self, file_data):
        """Parsear esquema de CSV"""
        import csv
        import io

        columns = []
        content = file_data.decode('utf-8')
        reader = csv.DictReader(io.StringIO(content))

        for field_name in reader.fieldnames or []:
            columns.append({
                'name': field_name,
                'type': 'varchar',
                'nullable': True,
                'is_pk': False,
                'is_fk': False,
            })

        return columns

    def _parse_excel_schema(self, file_data):
        """Parsear esquema de Excel"""
        try:
            import openpyxl
            import io

            wb = openpyxl.load_workbook(io.BytesIO(file_data), read_only=True)
            sheet = wb.active

            columns = []
            for cell in sheet[1]:
                if cell.value:
                    columns.append({
                        'name': str(cell.value),
                        'type': 'varchar',
                        'nullable': True,
                        'is_pk': False,
                        'is_fk': False,
                    })

            return columns

        except ImportError:
            _logger.warning('openpyxl no instalado, no se puede leer Excel')
            return []

    def _parse_json_schema(self, file_data):
        """Parsear esquema de JSON"""
        data = json.loads(file_data.decode('utf-8'))

        # Asumir que es un array de objetos
        if isinstance(data, list) and len(data) > 0:
            sample = data[0]
            if isinstance(sample, dict):
                columns = []
                for key, value in sample.items():
                    col_type = 'varchar'
                    if isinstance(value, int):
                        col_type = 'integer'
                    elif isinstance(value, float):
                        col_type = 'float'
                    elif isinstance(value, bool):
                        col_type = 'boolean'

                    columns.append({
                        'name': key,
                        'type': col_type,
                        'nullable': True,
                        'is_pk': False,
                        'is_fk': False,
                    })
                return columns

        return []

    def _read_mssql_schema(self, connection):
        """Leer esquema de SQL Server"""
        # Implementación similar a PostgreSQL
        _logger.warning('SQL Server schema reader no implementado completamente')
        return []

    def _read_oracle_schema(self, connection):
        """Leer esquema de Oracle"""
        # Implementación similar a PostgreSQL
        _logger.warning('Oracle schema reader no implementado completamente')
        return []
